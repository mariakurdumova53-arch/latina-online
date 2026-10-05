import os, csv, io, json
from functools import wraps
from datetime import datetime, timezone
from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_wtf import CSRFProtect
from werkzeug.security import generate_password_hash, check_password_hash
from dotenv import load_dotenv

load_dotenv()
db=SQLAlchemy()
csrf=CSRFProtect()

class User(db.Model):
    id=db.Column(db.Integer,primary_key=True)
    name=db.Column(db.String(160),nullable=False)
    email=db.Column(db.String(255),unique=True,nullable=False,index=True)
    group_name=db.Column(db.String(80),default="")
    password_hash=db.Column(db.String(255),nullable=False)
    role=db.Column(db.String(20),nullable=False,default="student")
    active=db.Column(db.Boolean,default=True)
    created_at=db.Column(db.DateTime,default=lambda: datetime.now(timezone.utc))
    attempts=db.relationship("Attempt",backref="user",lazy=True,cascade="all, delete-orphan")

class Question(db.Model):
    id=db.Column(db.Integer,primary_key=True)
    topic=db.Column(db.String(160),nullable=False,index=True)
    question=db.Column(db.Text,nullable=False)
    a=db.Column(db.Text,nullable=False); b=db.Column(db.Text,nullable=False)
    c=db.Column(db.Text,nullable=False); d=db.Column(db.Text,nullable=False)
    correct=db.Column(db.String(1),nullable=False)
    explanation=db.Column(db.Text,default="")
    active=db.Column(db.Boolean,default=True)
    created_at=db.Column(db.DateTime,default=lambda: datetime.now(timezone.utc))

class Attempt(db.Model):
    id=db.Column(db.Integer,primary_key=True)
    user_id=db.Column(db.Integer,db.ForeignKey("user.id"),nullable=False,index=True)
    topic=db.Column(db.String(160),nullable=False)
    score=db.Column(db.Integer,nullable=False)
    total=db.Column(db.Integer,nullable=False)
    percent=db.Column(db.Float,nullable=False)
    started_at=db.Column(db.DateTime,default=lambda: datetime.now(timezone.utc),index=True)

def login_required(f):
    @wraps(f)
    def w(*a,**kw):
        if not session.get("user_id"): return redirect(url_for("login",next=request.path))
        return f(*a,**kw)
    return w
def teacher_required(f):
    @wraps(f)
    def w(*a,**kw):
        u=current_user()
        if not u or u.role!="teacher": flash("Доступ только для преподавателя.","error"); return redirect(url_for("index"))
        return f(*a,**kw)
    return w
def current_user():
    uid=session.get("user_id")
    return db.session.get(User,uid) if uid else None

def create_app():
    app=Flask(__name__)
    app.config["SECRET_KEY"]=os.environ.get("SECRET_KEY","dev-change-me")
    app.config["SQLALCHEMY_DATABASE_URI"]=os.environ.get("DATABASE_URL","sqlite:///latina.db")
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"]=False
    app.config["SESSION_COOKIE_HTTPONLY"]=True
    app.config["SESSION_COOKIE_SAMESITE"]="Lax"
    app.config["SESSION_COOKIE_SECURE"]=os.environ.get("COOKIE_SECURE","0")=="1"
    db.init_app(app); csrf.init_app(app)
    with app.app_context():
        db.create_all()
        seed_questions()
        seed_admin()

    @app.context_processor
    def inject():
        return {"me":current_user()}

    @app.route("/")
    def index(): return render_template("index.html")
    @app.route("/register",methods=["GET","POST"])
    def register():
        if request.method=="POST":
            name=request.form.get("name","").strip(); email=request.form.get("email","").strip().lower()
            group=request.form.get("group_name","").strip(); password=request.form.get("password","")
            if not name or not email or len(password)<8: flash("Заполните поля. Пароль — минимум 8 символов.","error")
            elif User.query.filter_by(email=email).first(): flash("Этот email уже зарегистрирован.","error")
            else:
                u=User(name=name,email=email,group_name=group,password_hash=generate_password_hash(password),role="student")
                db.session.add(u); db.session.commit(); session["user_id"]=u.id
                return redirect(url_for("index"))
        return render_template("auth.html",mode="register")
    @app.route("/login",methods=["GET","POST"])
    def login():
        if request.method=="POST":
            email=request.form.get("email","").strip().lower(); password=request.form.get("password","")
            u=User.query.filter_by(email=email,active=True).first()
            if u and check_password_hash(u.password_hash,password):
                session.clear(); session["user_id"]=u.id
                return redirect(request.args.get("next") or url_for("index"))
            flash("Неверный email или пароль.","error")
        return render_template("auth.html",mode="login")
    @app.route("/logout")
    def logout(): session.clear(); return redirect(url_for("index"))

    @app.route("/api/topics")
    def api_topics():
        return jsonify([x[0] for x in db.session.query(Question.topic).filter_by(active=True).distinct().order_by(Question.topic).all()])

    @app.route("/test/<path:topic>")
    @login_required
    def test(topic):
        qs=Question.query.filter_by(active=True,topic=topic).all()
        return render_template("test.html",topic=topic,questions=[qdict(q,False) for q in qs])
    @app.route("/api/attempt",methods=["POST"])
    @login_required
    def save_attempt():
        data=request.get_json(force=True); answers=data.get("answers",{})
        topic=data.get("topic","")
        qs=Question.query.filter_by(active=True,topic=topic).all()
        score=sum(1 for q in qs if answers.get(str(q.id))==q.correct)
        a=Attempt(user_id=session["user_id"],topic=topic,score=score,total=len(qs),percent=round(score*100/len(qs),1) if qs else 0)
        db.session.add(a); db.session.commit()
        return jsonify(ok=True,score=score,total=len(qs),percent=a.percent)
    @app.route("/results")
    @login_required
    def results():
        return render_template("results.html",attempts=Attempt.query.filter_by(user_id=session["user_id"]).order_by(Attempt.id.desc()).all())

    @app.route("/admin")
    @teacher_required
    def admin():
        stats={"students":User.query.filter_by(role="student").count(),"attempts":Attempt.query.count(),
               "questions":Question.query.filter_by(active=True).count(),
               "average":round(db.session.query(db.func.avg(Attempt.percent)).scalar() or 0,1)}
        return render_template("admin.html",stats=stats,results=Attempt.query.order_by(Attempt.id.desc()).limit(300).all(),
                               questions=Question.query.order_by(Question.topic,Question.id).all())

    @app.route("/admin/question",methods=["POST"])
    @teacher_required
    def add_question():
        f=request.form
        q=Question(topic=f["topic"].strip(),question=f["question"].strip(),a=f["a"].strip(),b=f["b"].strip(),
                   c=f["c"].strip(),d=f["d"].strip(),correct=f["correct"],explanation=f.get("explanation","").strip())
        if q.correct not in "ABCD": flash("Правильный ответ должен быть A/B/C/D","error")
        else: db.session.add(q); db.session.commit(); flash("Вопрос добавлен.","ok")
        return redirect(url_for("admin"))
    @app.route("/admin/question/<int:id>/delete",methods=["POST"])
    @teacher_required
    def delete_question(id):
        q=db.session.get(Question,id)
        if q: q.active=False; db.session.commit()
        return redirect(url_for("admin"))

    @app.route("/admin/import",methods=["POST"])
    @teacher_required
    def import_excel():
        from openpyxl import load_workbook
        file=request.files.get("file")
        if not file: flash("Выберите Excel-файл.","error"); return redirect(url_for("admin"))
        try:
            wb=load_workbook(file,read_only=True,data_only=True)
            ws=wb["Тесты"] if "Тесты" in wb.sheetnames else wb.active
            added=0
            for r in ws.iter_rows(min_row=2,values_only=True):
                if not r[2]: continue
                correct=str(r[7] or "").strip().upper()
                if correct not in "ABCD": continue
                db.session.add(Question(topic=str(r[0] or "Общие"),question=str(r[2]),a=str(r[3] or ""),b=str(r[4] or ""),c=str(r[5] or ""),d=str(r[6] or ""),correct=correct))
                added+=1
            db.session.commit(); flash(f"Импортировано вопросов: {added}","ok")
        except Exception as e: flash("Ошибка импорта Excel: "+str(e),"error")
        return redirect(url_for("admin"))

    @app.route("/admin/export.csv")
    @teacher_required
    def export_csv():
        out=io.StringIO(); w=csv.writer(out); w.writerow(["topic","question","A","B","C","D","correct","explanation"])
        for q in Question.query.filter_by(active=True).order_by(Question.topic,Question.id): w.writerow([q.topic,q.question,q.a,q.b,q.c,q.d,q.correct,q.explanation])
        mem=io.BytesIO(out.getvalue().encode("utf-8-sig"))
        return send_file(mem,as_attachment=True,download_name="latina_questions.csv",mimetype="text/csv")

    return app

def qdict(q,include_answer=False):
    d={"id":q.id,"question":q.question,"a":q.a,"b":q.b,"c":q.c,"d":q.d}
    if include_answer: d["correct"]=q.correct
    return d

def seed_questions():
    if Question.query.count(): return
    path=os.path.join(os.path.dirname(os.path.dirname(__file__)),"questions.json")
    if not os.path.exists(path): return
    for q in json.load(open(path,encoding="utf8")):
        if q["correct"] in "ABCD":
            db.session.add(Question(**q))
    db.session.commit()

def seed_admin():
    email=os.environ.get("ADMIN_EMAIL"); password=os.environ.get("ADMIN_PASSWORD")
    if not email or not password or len(password)<12: return
    if not User.query.filter_by(email=email.lower()).first():
        db.session.add(User(name="Преподаватель",email=email.lower(),password_hash=generate_password_hash(password),role="teacher"))
        db.session.commit()
